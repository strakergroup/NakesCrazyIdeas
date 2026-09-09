import io

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas

from assurance.checkers import equivalent
from assurance.extraction import extract


@pytest.mark.parametrize(
    "actual,expected,unit",
    [
        ("24 months", "two years", None),
        ("2 kg", "2000 g", None),
        (24, "two years", "months"),
        ("1 m", "100 cm", None),
    ],
)
def test_defined_equivalent_quantity(actual, expected, unit):
    assert equivalent(actual, expected, unit)


@pytest.mark.parametrize(
    "actual,expected", [("30 days", "1 month"), ("18% consumption", "18% bills"), (True, 1), ("2 kg", "2 g")]
)
def test_undefined_or_false_equivalence_does_not_pass(actual, expected):
    assert not equivalent(actual, expected)


def test_plain_text_and_markdown_offsets():
    raw = "Café warranty.\n\n# Support\nContact us.".encode()
    result = extract(raw, "md")
    assert result["supported"]
    for i in result["items"]:
        loc = i["locator"]
        assert raw.decode()[loc["start"] : loc["end"]] == i["text"]
    assert result["gaps"][0]["modality"] == "rendered"


def test_product_json_values_and_mandated_text(h):
    brief, _, _ = h.setup(
        brief_extra={
            "required_fields": ["/name", "/warranty"],
            "structured_values": [{"path": "/warranty", "expected": "24 months"}],
            "required_text": ["Contact support"],
        }
    )
    sub = h.submit(
        brief, content={"name": "Device", "warranty": "two years", "help": "Contact support"}, fmt="json"
    )
    run = h.run(sub)
    checks = run["snapshot"]["checks"]
    assert checks["structured_values:0"]["outcome"] == "pass"
    assert checks["required_fields:0"]["outcome"] == "pass"
    assert checks["required_text:0"]["outcome"] == "pass"
    original = h.request("GET", f"/versions/{sub['version_id']}")
    assert {i["locator"]["path"] for i in original["payload"]["extraction"]["items"]} == {
        "/name",
        "/warranty",
        "/help",
    }


def test_known_missing_requirement_is_failure(h):
    brief, _, _ = h.setup(brief_extra={"required_fields": ["/missing"]})
    sub = h.submit(brief, content={"name": "Device"}, fmt="json")
    run = h.run(sub)
    assert run["snapshot"]["checks"]["required_fields:0"]["outcome"] == "fail"
    assert any(f["primary_criterion"] == "C2" for f in run["snapshot"]["findings"])


@pytest.mark.parametrize(
    "raw,fmt",
    [
        (b'{"x":1,"x":2}', "json"),
        (b'{"x":NaN}', "json"),
        (b"{bad}", "json"),
        (b"\xff\xfe", "txt"),
        (b"bad pdf", "pdf"),
        (b"bad docx", "docx"),
    ],
)
def test_failed_extraction_visible(raw, fmt):
    result = extract(raw, fmt)
    assert not result["supported"] and result["errors"]


def test_readable_pdf_and_scan_have_truthful_coverage():
    stream = io.BytesIO()
    canvas = Canvas(stream)
    canvas.drawString(50, 750, "Readable warranty information.")
    canvas.save()
    x = extract(stream.getvalue(), "pdf")
    assert x["supported"] and x["items"][0]["locator"]["page"] == 1
    assert x["gaps"]
    scan = PdfWriter()
    scan.add_blank_page(100, 100)
    b = io.BytesIO()
    scan.write(b)
    x = extract(b.getvalue(), "pdf")
    assert not x["supported"] and "PAGE_WITHOUT_TEXT:1" in x["errors"]


def test_docx_real_paragraph_and_table_extraction():
    doc = Document()
    doc.add_paragraph("Warranty")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Duration"
    table.cell(0, 1).text = "24 months"
    doc.sections[0].header.paragraphs[0].text = "Header"
    stream = io.BytesIO()
    doc.save(stream)
    x = extract(stream.getvalue(), "docx")
    assert x["supported"] and "24 months" in x["text"]
    assert any("/cell[" in i["locator"]["path"] for i in x["items"])
    assert "headers" in x["gaps"][0]["reason"]


def test_exclusions_cannot_hide_unsupported_visual_examinations(h):
    brief, _, _ = h.setup(
        brief_extra={
            "exclusions": {"U3": {"reason": "Not applicable", "scope": "test", "evidence": "owner decision"}}
        }
    )
    h.request(
        "POST",
        "/versions",
        {"brief_id": brief, "title": "image", "content": "![chart](chart.png)", "format": "md"},
        "submitter",
        422,
        headers={"Idempotency-Key": "visual"},
    )
